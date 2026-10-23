import os
from datetime import datetime
import requests
import re
from bs4 import BeautifulSoup
import openpyxl
import asyncio

# 日志记录函数
def log_message(user_id: str, message: str, log_type: str = "INFO"):
    # 获取当前进程的PID
    pid = os.getpid()
    user_folder = f"user_{user_id}_pid_{pid}"
    os.makedirs(f"logs/{user_folder}", exist_ok=True)

    # 创建日志文件路径
    log_file = f"logs/{user_folder}/log_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"

    # 格式化日志内容
    log_content = f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}][{log_type}] {message}\n"

    # 写入日志文件
    with open(log_file, "a", encoding="utf-8") as f:
        f.write(log_content)

    # 同时打印到控制台（可选）
    print(log_content.strip())

class Crawler:
    def __init__(self, url, headers, given_url):
        self.url = url
        self.headers = headers
        self.given_url = given_url
        self.content = ''

    async def get_web_content(self, websocket, user_id):
        url = f'{self.url}{self.given_url}'
        try:
            response = requests.get(url, headers=self.headers)
            response.raise_for_status()
            self.content = response.text
            await websocket.send_text(f"成功获取网页内容: {url}")
            log_message(user_id, f"成功获取网页内容: {url}", log_type="INFO")
            return self.content
        except requests.RequestException as e:
            await websocket.send_text(f"获取网页内容失败: {url}, 错误: {e}")
            log_message(user_id, f"获取网页内容失败: {url}, 错误: {e}", log_type="ERROR")
            self.content = ''
            return self.content

    def extract_image_urls(self):
        if self.content == "":
            print("请先运行get_web_content函数，或者抓取出错")
            return []
        else:
            image_urls = re.findall(r'(https?://[^\s]+(?:\.(?:jpg|jpeg|png|gif|bmp)))', self.content)
            print(image_urls)
            return image_urls

    @staticmethod
    async def crawl_data(industry_name, file_path, websocket, user_id):
        try:
            workbook = openpyxl.load_workbook(file_path)
            worksheet_name = f'materials/{industry_name}资料'
            total_text = ''
            if worksheet_name in workbook.sheetnames:
                worksheet = workbook[worksheet_name]
                nrows = worksheet.max_row
                for idx in range(2, nrows + 1):
                    url_link = worksheet.cell(idx, 2).value
                    if url_link is not None:
                        await websocket.send_text(f"开始抓取网页: {url_link}")
                        try:
                            text = await Crawler.scrape_and_update_batch(url_link, websocket, user_id)
                            if not text:
                                text = "无内容"
                            total_text = total_text + '\n' + text
                        except Exception as e:
                            await websocket.send_text(f"抓取网页时出错: {url_link}, 错误: {e}")
                            log_message(user_id, f"抓取网页时出错: {url_link}, 错误: {e}", log_type="ERROR")
            await websocket.send_text("辅助数据源抓取完成")
            log_message(user_id, "辅助数据源抓取完成", log_type="INFO")
            return total_text
        except Exception as e:
            await websocket.send_text(f"读取Excel文件时出错: {e}")
            log_message(user_id, f"读取Excel文件时出错: {e}", log_type="ERROR")
            return ""

    @staticmethod
    async def crawl_url_data(url, websocket, user_id):
        try:
            text = await Crawler.scrape_and_update_batch(url, websocket, user_id)
            return text
        except Exception as e:
            await websocket.send_text(f"抓取网页数据时出错: {url}, 错误: {e}")
            log_message(user_id, f"抓取网页数据时出错: {url}, 错误: {e}", log_type="ERROR")
            return None

    @staticmethod
    async def scrape_and_update_batch(url, websocket, user_id):
        try:
            response = requests.get(url)
            response.raise_for_status()
            soup = BeautifulSoup(response.text, 'html.parser')
            text = soup.get_text()
            await websocket.send_text(f"成功抓取网页内容: {url}")
            log_message(user_id, f"成功抓取网页内容: {url}", log_type="INFO")
            return text
        except requests.RequestException as e:
            await websocket.send_text(f"抓取网页时出错: {url}, 错误: {e}")
            log_message(user_id, f"抓取网页时出错: {url}, 错误: {e}", log_type="ERROR")
            return ""