import mysql.connector
import json
from openai import OpenAI
from openpyxl import load_workbook, Workbook
import os

# 数据库连接配置
db_config = {
    'user': 'remote_user',
    'password': os.getenv('DB_PASSWORD_REMOTE'),
    'host': '8.141.23.4',
    'database': 'business_data'
}

# 初始化OpenAI客户端
client = OpenAI(
    api_key=os.getenv("VOLCENGINE_ARK_API_KEY"),
    base_url="https://ark.cn-beijing.volces.com/api/v3"
)

# 连接到数据库
def connect_to_db(config):
    try:
        conn = mysql.connector.connect(**config)
        return conn
    except mysql.connector.Error as err:
        print(f"Error: {err}")
        return None

# 从数据库中提取数据
def fetch_data(conn):
    cursor = conn.cursor(dictionary=True)
    
    # 提取潜在客户数据
    cursor.execute("SELECT company_name, customer_company, buy_product FROM potential_company_0214 LIMIT 50")
    potential_companies = cursor.fetchall()
    
    # 提取供应商数据
    cursor.execute("SELECT combined_labels FROM cleaned_supplier_labels0214")
    suppliers = cursor.fetchall()
    
    cursor.close()
    return potential_companies, suppliers

# 构建提示词
def build_prompt(potential_company, suppliers):
    prompt = {
        "潜在客户信息": [potential_company],  # 只处理一个潜在客户
        "供应商信息": suppliers,
        "匹配要求": """
        根据潜在客户的采购需求（buy_product）和供应商的产品标签（combined_labels），进行语义匹配。
        匹配规则：
        1. 如果采购需求中包含品牌名称或公司名称，则优先匹配该品牌或公司的供应商。
        2. 如果没有明确品牌匹配，则分析产品名称的近义或相关性，进行匹配。
        输出结果应包含：
        - 每个潜在客户的公司名称（company_name），采购需求（buy_product），供应商的产品标签（combined_labels），匹配到的供应商品牌名称以及匹配原因，请你以JSON格式返回结果：
          [
            {
              "公司名称": "南京圣卡孚科技有限公司",
              "采购需求": "PLC模块",
              "供应商产品标签": "西门子_工控_PLC_其他PLC模块",
              "供应商品牌名称": "戴思乐",
              "匹配原因": "采购需求中包含'PLC模块'，与供应商的'其他PLC模块'语义相关。"
            }
          ]
        """
    }
    return json.dumps(prompt, ensure_ascii=False)

# 调用大模型接口
def call_model_api(prompt):
    try:
        response = client.chat.completions.create(
            model="ep-20250205154622-7gm5g",
            messages=[
                {'role': 'system', 'content': prompt},
                {'role': 'user', 'content': "请根据上述信息进行供应商和潜在客户的匹配。"}
            ],
            response_format={"type": "json_object"}
        )
        return response.choices[0].message.content
    except Exception as e:
        print(f"API调用失败：{e}")
        return None

# 尝试修复不完整的JSON数据
def fix_incomplete_json(data):
    try:
        # 尝试直接解析JSON数据
        return json.loads(data)
    except json.JSONDecodeError as e:
        # 如果解析失败，尝试修复不完整的JSON数据
        print(f"JSON解析失败，尝试修复不完整的JSON数据：{e}")
        # 尝试移除尾部的不完整部分
        data = data[:data.rfind('}')] + "}"
        return json.loads(data)

# 将结果写入Excel文件
def write_to_excel(data, file_name="potential_customers.xlsx"):
    if not os.path.exists(file_name):  # 如果文件不存在，新建文件并写入表头
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["公司名称", "采购需求", "供应商产品标签", "供应商品牌名称", "匹配原因"])
    else:  # 如果文件已存在，加载文件
        workbook = load_workbook(file_name)
        sheet = workbook.active
    
    try:
        # 尝试修复并解析JSON数据
        results = fix_incomplete_json(data)
        if not isinstance(results, list):  # 确保结果是列表格式
            raise ValueError("返回的JSON数据不是列表格式")
        for result in results:  # 遍历列表中的每个匹配结果
            if not isinstance(result, dict):  # 确保每个结果是字典格式
                raise ValueError("列表中的元素不是字典格式")
            sheet.append([
                result.get("公司名称", ""),
                result.get("采购需求", ""),
                result.get("供应商产品标签", ""),
                result.get("供应商品牌名称", ""),
                result.get("匹配原因", "")
            ])
    except (json.JSONDecodeError, ValueError) as e:
        print(f"数据格式错误：{e}")
        # 记录错误信息到日志文件
        with open("error_log.txt", "a") as log_file:
            log_file.write(f"数据格式错误：{e}\n原始数据：{data}\n\n")

    # 保存Excel文件
    workbook.save(file_name)
    print(f"结果已写入文件：{file_name}")

# 主程序
def main():
    conn = connect_to_db(db_config)
    if conn:
        potential_companies, suppliers = fetch_data(conn)
        for potential_company in potential_companies:
            prompt = build_prompt(potential_company, suppliers)
            result = call_model_api(prompt)
            if result:
                print("匹配结果：", result)
                write_to_excel(result)  # 将结果写入Excel
            else:
                print(f"未获取到匹配结果（公司名称：{potential_company['company_name']}），请检查API调用或数据格式。")
        conn.close()
    else:
        print("数据库连接失败，无法继续执行。")

if __name__ == "__main__":
    main()