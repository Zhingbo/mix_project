import requests
import re
from bs4 import BeautifulSoup
import openpyxl
from datetime import datetime
import os


class Crawler:
    def __init__(self, url, headers, given_url):
        self.url = url
        self.headers = headers
        self.given_url = given_url
        self.content = ''

    def get_web_content(self):
        url = f'{self.url}{self.given_url}'
        try:
            response = requests.get(url, headers=self.headers)
            response.raise_for_status()
            self.content = response.text
            print(f"成功获取网页内容: {url}")
            return self.content
        except requests.RequestException as e:
            print(f"获取网页内容失败: {url}, 错误: {e}")
            self.content = ''
            return self.content

    def extract_image_urls(self):
        if self.content == "":
            print("请先运行get_web_content函数，或者抓取出错")
            return []
        else:
            image_urls = re.findall(r'(https?://[^\s]+(?:\.(?:jpg|jpeg|png|gif|bmp)))', self.content)
            print(f"提取到的图片链接: {image_urls}")
            return image_urls

    @staticmethod
    def crawl_data(industry_name, file_path):
        try:
            workbook = openpyxl.load_workbook(file_path)
            worksheet_name = f'materials/{industry_name}资料'
            total_text = ''
            if worksheet_name in workbook.sheetnames:
                worksheet = workbook[worksheet_name]
                nrows = worksheet.max_row
                for idx in range(2, nrows + 1):
                    url_link = worksheet.cell(idx, 2).value
                    if url_link is not None:
                        print(f"抓取网页: {url_link}")
                        text = Crawler.scrape_and_update_batch(url_link)
                        if not text:
                            text = "无内容"
                        total_text = total_text + '\n' + text
            print(f"辅助数据源抓取完成，内容：{total_text}")
            return total_text
        except Exception as e:
            print(f"读取Excel文件时出错: {e}")
            return ""

    @staticmethod
    def crawl_url_data(url):
        text = Crawler.scrape_and_update_batch(url)
        return text

    @staticmethod
    def scrape_and_update_batch(url):
        try:
            response = requests.get(url)
            response.raise_for_status()
            soup = BeautifulSoup(response.text, 'html.parser')
            text = soup.get_text()
            print(f"成功抓取网页内容: {url}")
            return text
        except requests.RequestException as e:
            print(f"抓取网页时出错: {url}, 错误: {e}")
            return ""