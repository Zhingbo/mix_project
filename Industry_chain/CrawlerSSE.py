import os
from datetime import datetime
import aiohttp
import re
from bs4 import BeautifulSoup
import openpyxl

def log_message(user_id: str, message: str, log_type: str = "INFO"):
    # 获取当前进程的PID
    pid = os.getpid()
    user_folder = f"user_{user_id}_pid_{pid}"
    os.makedirs(f"logs/{user_folder}", exist_ok=True)

    # 创建日志文件路径
    log_file = f"logs/{user_folder}/log_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"

    # 格式化日志内容
    log_content = f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}][{log_type}] {message}\n"

    # 写入日志文件
    with open(log_file, "a", encoding="utf-8") as f:
        f.write(log_content)

    # 同时打印到控制台（可选）
    print(log_content.strip())

class Crawler:
    def __init__(self, url, headers, given_url):
        self.url = url
        self.headers = headers
        self.given_url = given_url
        self.content = ''

    async def get_web_content(self, user_id: str, send_message):
        url = f'{self.url}{self.given_url}'
        try:
            async with aiohttp.ClientSession() as session:
                async with session.get(url, headers=self.headers) as response:
                    response.raise_for_status()
                    self.content = await response.text()
                    log_message(user_id, f"成功获取网页内容: {url}", log_type="INFO")
                    await send_message(f"成功获取网页内容: {url}")
                    return f"成功获取网页内容: {url}"
        except aiohttp.ClientError as e:
            log_message(user_id, f"获取网页内容失败: {url}, 错误: {e}", log_type="ERROR")
            self.content = ''
            await send_message(f"获取网页内容失败: {url}, 错误: {e}")
            return f"获取网页内容失败: {url}, 错误: {e}"

    def extract_image_urls(self):
        if self.content == "":
            print("请先运行get_web_content函数，或者抓取出错")
            return []
        else:
            image_urls = re.findall(r'(https?://[^\s]+(?:\.(?:jpg|jpeg|png|gif|bmp)))', self.content)
            print(image_urls)
            return image_urls

    @staticmethod
    async def crawl_data(industry_name: str, file_path: str, user_id: str, send_message):
        try:
            workbook = openpyxl.load_workbook(file_path)
            worksheet_name = f'materials/{industry_name}资料'
            total_text = ''
            if worksheet_name in workbook.sheetnames:
                worksheet = workbook[worksheet_name]
                nrows = worksheet.max_row
                for idx in range(2, nrows + 1):
                    url_link = worksheet.cell(idx, 2).value
                    if url_link is not None:
                        log_message(user_id, f"开始抓取网页: {url_link}", log_type="INFO")
                        await send_message(f"开始抓取网页: {url_link}")
                        try:
                            text = await Crawler.crawl_url_data(url_link, user_id, send_message)
                            if not text:
                                text = "无内容"
                            total_text = total_text + '\n' + text
                        except Exception as e:
                            log_message(user_id, f"抓取网页时出错: {url_link}, 错误: {e}", log_type="ERROR")
                            await send_message(f"抓取网页时出错: {url_link}, 错误: {e}")
            log_message(user_id, "辅助数据源抓取完成", log_type="INFO")
            await send_message("辅助数据源抓取完成")
            return total_text
        except Exception as e:
            log_message(user_id, f"读取Excel文件时出错: {e}", log_type="ERROR")
            await send_message(f"读取Excel文件时出错: {e}")
            return f"读取Excel文件时出错: {e}"

    @staticmethod
    async def crawl_url_data(url: str, user_id: str, send_message):
        try:
            text = await Crawler.scrape_and_update_batch(url, user_id, send_message)
            return text
        except Exception as e:
            log_message(user_id, f"抓取网页数据时出错: {url}, 错误: {e}", log_type="ERROR")
            await send_message(f"抓取网页数据时出错: {url}, 错误: {e}")
            return f"抓取网页数据时出错: {url}, 错误: {e}"

    @staticmethod
    async def scrape_and_update_batch(url: str, user_id: str, send_message):
        try:
            async with aiohttp.ClientSession() as session:
                async with session.get(url) as response:
                    response.raise_for_status()
                    text = await response.text()
                    log_message(user_id, f"成功抓取网页内容: {url}", log_type="INFO")
                    await send_message(f"成功抓取网页内容: {url}")
                    return text
        except aiohttp.ClientError as e:
            log_message(user_id, f"抓取网页时出错: {url}, 错误: {e}", log_type="ERROR")
            await send_message(f"抓取网页时出错: {url}, 错误: {e}")
            return f"抓取网页时出错: {url}, 错误: {e}"