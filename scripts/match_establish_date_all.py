# -*- coding: utf-8 -*-
# 对 产业链导出_20261008 下每个 *_企业名单.xlsx 匹配成立时间, 输出 *_含成立时间.xlsx
import openpyxl
import pymysql
import os
import glob

DIR = "/root/mix_project/产业链导出_20261008"
SHEET_NAME = "企业表"
BATCH_SIZE = 500

ALIYUN_DB = {
    "host": "rm-2zej668f34v6xr91wdo.mysql.rds.aliyuncs.com",
    "port": 3306,
    "user": "root",
    "password": os.getenv("DB_PASSWORD_ZS"),
    "database": "zs_digital_staff_db",
    "charset": "utf8mb4",
    "connect_timeout": 60,
    "read_timeout": 60,
}
TABLE = "enterprise_establish_time_chain_scope_20260928"


def process_one(input_xlsx):
    output_xlsx = input_xlsx.replace("_企业名单.xlsx", "_企业名单_含成立时间.xlsx")
    print(f"\n{'='*60}\n处理: {os.path.basename(input_xlsx)}")

    wb_in = openpyxl.load_workbook(input_xlsx, read_only=True)
    ws = wb_in[SHEET_NAME]
    rows = []
    header = None
    for i, row in enumerate(ws.iter_rows(min_row=1, values_only=True)):
        if header is None:
            header = list(row)
            continue
        company = str(row[3]).strip() if row[3] else ""
        rows.append({"c1": row[0], "c2": row[1], "c3": row[2], "company": company})
    wb_in.close()
    print(f"  读到企业数: {len(rows)}")

    conn = pymysql.connect(**ALIYUN_DB)
    cursor = conn.cursor()
    result_map = {}
    unique_companies = list(dict.fromkeys([r["company"] for r in rows if r["company"]]))
    for bi in range(0, len(unique_companies), BATCH_SIZE):
        batch = unique_companies[bi:bi + BATCH_SIZE]
        ph = ",".join(["%s"] * len(batch))
        sql = f"SELECT company_name, establish_time, establish_year FROM {TABLE} WHERE company_name IN ({ph})"
        try:
            cursor.execute(sql, batch)
            for name, inc_time, reg_year in cursor.fetchall():
                result_map[name] = (inc_time, reg_year)
        except Exception as e:
            print(f"  批 {bi//BATCH_SIZE+1} 失败: {e}")
    cursor.close()
    conn.close()
    print(f"  匹配到成立时间: {len(result_map)} / {len(unique_companies)}")

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = SHEET_NAME
    ws_out.append(header + ["成立时间", "注册年份", "匹配状态"])
    matched_cnt = 0
    for r in rows:
        inc_time, reg_year = result_map.get(r["company"], (None, None))
        status = "已匹配" if (inc_time or reg_year) else "未匹配"
        if inc_time or reg_year:
            matched_cnt += 1
        if inc_time is None:
            inc_str = ""
        elif hasattr(inc_time, "strftime"):
            inc_str = inc_time.strftime("%Y-%m-%d")
        else:
            inc_str = str(inc_time).strip()[:10]
        ws_out.append([r["c1"], r["c2"], r["c3"], r["company"], inc_str, reg_year or "", status])
    wb_out.save(output_xlsx)
    print(f"  输出: {output_xlsx} (匹配 {matched_cnt}, 未匹配 {len(rows)-matched_cnt})")


def main():
    files = sorted(glob.glob(os.path.join(DIR, "*_企业名单.xlsx")))
    files = [f for f in files if "含成立时间" not in f]
    print(f"共 {len(files)} 个文件待处理:")
    for f in files:
        print("  ", os.path.basename(f))
    for f in files:
        process_one(f)
    print("\n全部完成")


if __name__ == "__main__":
    main()